from flask import render_template, request, send_file, redirect, url_for, flash
from datetime import datetime
from io import BytesIO
from urllib.parse import urlencode
import re

from openpyxl import Workbook

from database.connection import get_connection
from ofs.client import OFSClient
from core.auth import login_required, perm_required
from core.utils import xlsx_auto_width

def normalize_appt_number(value):
    s = str(value or "").strip()
    if not s:
        return ""
    return re.sub(r"-[^/]+(?=/)", "", s)
def init_app(app):

    @app.route("/ofs/atividades-base", methods=["GET"])
    @login_required
    @perm_required("ofs.atividades_base")
    def ofs_atividades_base():
        date_from = (request.args.get("dateFrom") or "").strip()
        date_to = (request.args.get("dateTo") or "").strip()
        resources = (request.args.get("resources") or "02").strip()

        selected_types = [t.strip() for t in request.args.getlist("activityType") if (t or "").strip()]
        selected_statuses = [s.strip() for s in request.args.getlist("status") if (s or "").strip()]

        per_page = 100
        page = request.args.get("page", default=1, type=int)
        if page < 1:
            page = 1
        offset = (page - 1) * per_page

        conn = get_connection()
        cur = conn.cursor(dictionary=True)

        try:
            cur.execute("""
                SELECT DISTINCT activity_type
                FROM ofs_atividades_base
                WHERE activity_type IS NOT NULL AND activity_type <> ''
                ORDER BY activity_type
            """)
            activity_types = [r["activity_type"] for r in cur.fetchall()]

            cur.execute("""
                SELECT DISTINCT origin_bucket
                FROM ofs_atividades_base
                WHERE origin_bucket IS NOT NULL AND origin_bucket <> ''
                ORDER BY origin_bucket
            """)
            buckets = [r["origin_bucket"] for r in cur.fetchall()]

            cur.execute("""
                SELECT DISTINCT status
                FROM ofs_atividades_base
                WHERE status IS NOT NULL AND status <> ''
                ORDER BY status
            """)
            statuses = [r["status"] for r in cur.fetchall()]

            where = ["1=1"]
            params = []

            if date_from and date_to:
                where.append("b.`date` BETWEEN %s AND %s")
                params.extend([date_from, date_to])
            elif date_from:
                where.append("b.`date` >= %s")
                params.append(date_from)
            elif date_to:
                where.append("b.`date` <= %s")
                params.append(date_to)

            if selected_types:
                placeholders = ",".join(["%s"] * len(selected_types))
                where.append(f"b.activity_type IN ({placeholders})")
                params.extend(selected_types)

            if selected_statuses:
                placeholders = ",".join(["%s"] * len(selected_statuses))
                where.append(f"b.status IN ({placeholders})")
                params.extend(selected_statuses)

            where_sql = " AND ".join(where)

            cur.execute(f"""
                SELECT COUNT(*) AS total
                FROM ofs_atividades_base b
                WHERE {where_sql}
            """, tuple(params))
            total = int((cur.fetchone() or {}).get("total") or 0)

            total_pages = max(1, (total + per_page - 1) // per_page)
            if page > total_pages:
                page = total_pages
                offset = (page - 1) * per_page

            cur.execute(f"""
                SELECT
                    b.activity_id AS activityId,
                    b.city,
                    b.activity_type AS activityType,
                    COALESCE(atm.label_pt, b.activity_type) AS activityType_pt,
                    b.appt_number AS apptNumber,
                    b.origin_bucket AS XA_ORIGIN_BUCKET,
                    b.resource_id AS resourceId,
                    b.status AS status,
                    COALESCE(sm.label_pt, b.status) AS status_pt,
                    b.xa_org_sys AS XA_ORG_SYS,
                    b.`date` AS date
                FROM ofs_atividades_base b
                LEFT JOIN ofs_activity_type_map atm
                  ON atm.code = b.activity_type AND atm.is_active = 1
                LEFT JOIN ofs_status_map sm
                  ON sm.code = b.status AND sm.is_active = 1
                WHERE {where_sql}
                ORDER BY b.`date` ASC, b.activity_id ASC
                LIMIT %s OFFSET %s
            """, tuple(params + [per_page, offset]))
            items = cur.fetchall()

        finally:
            cur.close()
            conn.close()

        return render_template(
            "atividades_base/ofs_atividades_base.html",
            items=items,
            date_from=date_from,
            date_to=date_to,
            resources=resources,
            total=total,
            page=page,
            total_pages=total_pages,
            per_page=per_page,
            activity_types=activity_types,
            buckets=buckets,
            statuses=statuses,
            selected_types=selected_types,
            selected_statuses=selected_statuses,
        )

    @app.route("/ofs/atividades-base/exportar", methods=["POST"])
    @login_required
    @perm_required("ofs.atividades_base")
    def ofs_atividades_base_exportar():
        activity_types = [t.strip() for t in request.form.getlist("activityType") if (t or "").strip()]
        statuses = [s.strip() for s in request.form.getlist("status") if (s or "").strip()]
        date_from = (request.form.get("dateFrom") or "").strip()
        date_to = (request.form.get("dateTo") or "").strip()
        buckets = [b.strip() for b in request.form.getlist("buckets") if (b or "").strip()]

        try:
            dt_from = datetime.strptime(date_from, "%Y-%m-%d").date()
            dt_to = datetime.strptime(date_to, "%Y-%m-%d").date()
        except Exception:
            flash("Informe um período válido (De / Até).", "danger")
            return redirect(url_for("ofs_atividades_base"))

        if dt_to < dt_from:
            flash("O campo 'Até' não pode ser menor que 'De'.", "danger")
            return redirect(url_for("ofs_atividades_base", dateFrom=date_from, dateTo=date_to))

        conn = get_connection()
        cur = conn.cursor(dictionary=True)

        try:
            sql = """
                SELECT
                    b.activity_id AS activityId,
                    b.city,
                    b.activity_type AS activityType,
                    COALESCE(atm.label_pt, b.activity_type) AS activityType_pt,
                    b.appt_number AS apptNumber,
                    b.origin_bucket AS XA_ORIGIN_BUCKET,
                    b.resource_id AS resourceId,
                    b.status AS status,
                    COALESCE(sm.label_pt, b.status) AS status_pt,
                    b.xa_org_sys AS XA_ORG_SYS,
                    b.`date` AS date
                FROM ofs_atividades_base b
                LEFT JOIN ofs_activity_type_map atm
                  ON atm.code = b.activity_type AND atm.is_active = 1
                LEFT JOIN ofs_status_map sm
                  ON sm.code = b.status AND sm.is_active = 1
                WHERE b.`date` BETWEEN %s AND %s
            """
            params = [date_from, date_to]

            if activity_types:
                placeholders = ",".join(["%s"] * len(activity_types))
                sql += f" AND activity_type IN ({placeholders})"
                params.extend(activity_types)

            if buckets:
                placeholders = ",".join(["%s"] * len(buckets))
                sql += f" AND origin_bucket IN ({placeholders})"
                params.extend(buckets)

            if statuses:
                placeholders = ",".join(["%s"] * len(statuses))
                sql += f" AND status IN ({placeholders})"
                params.extend(statuses)

            sql += " ORDER BY `date` ASC, last_seen_at DESC"

            cur.execute(sql, tuple(params))
            rows = cur.fetchall()

        finally:
            cur.close()
            conn.close()

        wb = Workbook()
        ws = wb.active
        ws.title = "AtividadesBase"

        headers = [
            "activityId",
            "city",
            "activityType",
            "activityType_pt",
            "apptNumber",
            "XA_ORIGIN_BUCKET",
            "resourceId",
            "status",
            "status_pt",
            "XA_ORG_SYS",
            "date",
            "last_seen_at",
        ]
        ws.append(headers)

        for r in rows:
            ws.append([r.get(h) for h in headers])

        xlsx_auto_width(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"ofs_atividades_base_{dt_from}_{dt_to}_{stamp}.xlsx"

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    @app.route("/ofs/atividades-base/importar", methods=["POST"])
    @login_required
    @perm_required("ofs.atividades_base")
    def ofs_atividades_base_importar():
        today = datetime.now().strftime("%Y-%m-%d")
        date_from = (request.form.get("dateFrom") or today).strip()
        date_to = (request.form.get("dateTo") or today).strip()
        resources = (request.form.get("resources") or "02").strip()

        try:
            dt_from = datetime.strptime(date_from, "%Y-%m-%d").date()
            dt_to = datetime.strptime(date_to, "%Y-%m-%d").date()
            if dt_to < dt_from:
                flash("O campo 'Até' não pode ser menor que 'De'.", "danger")
                return redirect(url_for("ofs_atividades_base", dateFrom=date_from, dateTo=date_to, resources=resources))
        except Exception:
            flash("Informe um período válido (De / Até).", "danger")
            return redirect(url_for("ofs_atividades_base", dateFrom=date_from, dateTo=date_to, resources=resources))

        client = OFSClient()

        fields = [
            "activityId",
            "city",
            "activityType",
            "apptNumber",
            "XA_ORIGIN_BUCKET",
            "resourceId",
            "status",
            "XA_ORG_SYS",
            "date",
        ]

        base_params = {
            "dateFrom": date_from,
            "dateTo": date_to,
            "resources": resources,
            "fields": ",".join(fields),
            "limit": 2000,
            "offset": 0,
        }

        items = []
        has_more = True
        max_pages = 30
        page = 0

        try:
            while has_more and page < max_pages:
                qs = urlencode(base_params, safe="=,'")
                url = f"{client.base_url}/activities/?{qs}"
                data = client.authenticated_get(url)

                batch = data.get("items") or []
                items.extend(batch)

                has_more = bool(data.get("hasMore"))
                if has_more:
                    base_params["offset"] = len(items)
                page += 1

        except Exception as e:
            flash(f"❌ Falha ao importar da API: {e}", "danger")
            return redirect(url_for("ofs_atividades_base", dateFrom=date_from, dateTo=date_to, resources=resources))

        conn = get_connection()
        cur = conn.cursor()

        inserted = 0
        updated = 0

        sql = """
            INSERT INTO ofs_atividades_base
            (
                activity_id,
                city,
                activity_type,
                appt_number,
                appt_number_norm,
                origin_bucket,
                resource_id,
                status,
                xa_org_sys,
                `date`
            )
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            ON DUPLICATE KEY UPDATE
                city = VALUES(city),
                activity_type = VALUES(activity_type),
                appt_number = VALUES(appt_number),
                appt_number_norm = VALUES(appt_number_norm),
                origin_bucket = VALUES(origin_bucket),
                resource_id = VALUES(resource_id),
                status = VALUES(status),
                xa_org_sys = VALUES(xa_org_sys),
                `date` = VALUES(`date`),
                last_seen_at = NOW()
        """
        for a in items:
            activity_id = str(a.get("activityId") or "").strip()
            if not activity_id:
                continue

            appt_number = str(a.get("apptNumber") or "").strip() or None
            appt_number_norm = normalize_appt_number(appt_number)

            cur.execute(sql, (
                activity_id,
                str(a.get("city") or "") or None,
                str(a.get("activityType") or "") or None,
                appt_number,
                appt_number_norm or None,
                str(a.get("XA_ORIGIN_BUCKET") or "") or None,
                str(a.get("resourceId") or "") or None,
                str(a.get("status") or "") or None,
                str(a.get("XA_ORG_SYS") or "") or None,
                str(a.get("date") or "") or None,
            ))

            if cur.rowcount == 1:
                inserted += 1
            elif cur.rowcount == 2:
                updated += 1

        conn.commit()
        cur.close()
        conn.close()

        flash(f"Importação concluída. Novos: {inserted} | Atualizados: {updated} | Total API: {len(items)}", "success")
        return redirect(url_for("ofs_atividades_base", dateFrom=date_from, dateTo=date_to, resources=resources))